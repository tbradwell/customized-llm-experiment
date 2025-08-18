import sys
sys.path.append('.')
from src.evaluation.metrics import MetricsCalculator, COMETEvaluator
from src.evaluation.llm_judge import LLMJudge
from src.core.document_processor import DocumentProcessor
import json
from docx import Document

print('=== EVALUATING BOTH APPROACHES ===')

# Initialize evaluators
metrics_calc = MetricsCalculator()
comet_eval = COMETEvaluator()
llm_judge = LLMJudge()
doc_processor = DocumentProcessor()

# Load ground truth
gt_doc = Document('examples/amit_test/gt.docx')
gt_text = doc_processor.extract_text_content(gt_doc)
reference_contracts = [gt_text]

print(f'Ground truth loaded: {len(gt_text)} characters')

# Load both generated contracts
from src.core.content_generator import IntelligentContentGenerator, GenerationContext

generator = IntelligentContentGenerator()
skeleton_doc = Document('examples/amit_test/sekeleton_oracle.docx')
skeleton_text = doc_processor.extract_text_content(skeleton_doc)
placeholders = [p.field_name for p in doc_processor.find_placeholders(skeleton_doc)]

context = GenerationContext(
    contract_type='legal_claim',
    skeleton_text=skeleton_text,
    placeholders=placeholders,
    contract_data={
        'contract_type': 'legal_claim',
        'client_name': 'מבקש התביעה',
        'provider_name': 'נתבעת',
        'source_content': 'תוכן מקור'
    }
)

print('\nGenerating contracts for evaluation...')

# Generate without iteration
print('Generating without iteration...')
result_old = generator.generate_contract_content(context)
contract_old = result_old.generated_content

# Generate with iteration  
print('Generating with iteration...')
result_new = generator.generate_complete_contract(context)
contract_new = result_new.generated_content

print(f'Old method: {len(contract_old)} chars')
print(f'New method: {len(contract_new)} chars')

# Evaluate both approaches
print('\n=== EVALUATING WITHOUT ITERATION ===')
eval_old = {}

print('Running BLEU...')
bleu_old = metrics_calc.calculate_bleu_score(contract_old, reference_contracts)
eval_old['bleu'] = bleu_old.score
print(f'BLEU: {bleu_old.score:.4f}')

print('Running ROUGE...')
rouge_old = metrics_calc.calculate_rouge_scores(contract_old, reference_contracts)
eval_old['rouge'] = rouge_old.score
print(f'ROUGE: {rouge_old.score:.4f}')

print('Running COMET...')
comet_old = comet_eval.calculate_comet_score(contract_old, reference_contracts, skeleton_text)
eval_old['comet'] = comet_old.score
print(f'COMET: {comet_old.score:.4f}')

print('Running LLM Judge...')
llm_old = llm_judge.evaluate_contract(contract_old, context.contract_data, reference_contracts)
eval_old['llm_judge'] = llm_old.score
print(f'LLM Judge: {llm_old.score:.4f}')

print('\n=== EVALUATING WITH ITERATION ===')
eval_new = {}

print('Running BLEU...')
bleu_new = metrics_calc.calculate_bleu_score(contract_new, reference_contracts)
eval_new['bleu'] = bleu_new.score
print(f'BLEU: {bleu_new.score:.4f}')

print('Running ROUGE...')
rouge_new = metrics_calc.calculate_rouge_scores(contract_new, reference_contracts)
eval_new['rouge'] = rouge_new.score
print(f'ROUGE: {rouge_new.score:.4f}')

print('Running COMET...')
comet_new = comet_eval.calculate_comet_score(contract_new, reference_contracts, skeleton_text)
eval_new['comet'] = comet_new.score
print(f'COMET: {comet_new.score:.4f}')

print('Running LLM Judge...')
llm_new = llm_judge.evaluate_contract(contract_new, context.contract_data, reference_contracts)
eval_new['llm_judge'] = llm_new.score
print(f'LLM Judge: {llm_new.score:.4f}')

# Summary comparison
print('\n=== METRIC COMPARISON SUMMARY ===')
print('Metric          | Without Iter | With Iter   | Improvement')
print('----------------|--------------|-------------|------------')
for metric in eval_old.keys():
    old_val = eval_old[metric]
    new_val = eval_new[metric]
    improvement = new_val - old_val
    improvement_pct = (improvement / old_val * 100) if old_val != 0 else 0
    print(f'{metric:15} | {old_val:11.4f} | {new_val:10.4f} | {improvement:+.4f} ({improvement_pct:+.1f}%)')

# Show iteration details
print('\n=== ITERATION DETAILS ===')
if result_new.generation_metadata:
    meta = result_new.generation_metadata
    print(f'Placeholder completion iterations: {meta.get("completion_iterations", "N/A")}')
    print(f'Initial placeholders found: {meta.get("initial_placeholders", "N/A")}')
    print(f'Final placeholders remaining: {meta.get("final_placeholders", "N/A")}')
    print(f'Completion success: {meta.get("completion_success", "N/A")}')
    print(f'Refinement iterations: {meta.get("refinement_iterations", "N/A")}')
    print(f'Refinement complete: {meta.get("refinement_complete", "N/A")}')
    print(f'Final LLM decision: {meta.get("final_decision", "N/A")}')

# Save results
comparison_results = {
    'without_iteration': {
        'content_length': len(contract_old),
        'metrics': eval_old
    },
    'with_iteration': {
        'content_length': len(contract_new),
        'metrics': eval_new,
        'iteration_details': result_new.generation_metadata
    },
    'improvements': {
        metric: eval_new[metric] - eval_old[metric] 
        for metric in eval_old.keys()
    }
}

with open('examples/amit_test_output/evaluation_comparison.json', 'w', encoding='utf-8') as f:
    json.dump(comparison_results, f, indent=2, ensure_ascii=False)

print('\nDetailed comparison saved to evaluation_comparison.json')